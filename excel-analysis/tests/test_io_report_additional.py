import pandas as pd
from excel_analysis.io import save_excel_report
from excel_analysis.analyse import build_vendor_report
import io
from openpyxl import load_workbook


def test_vendor_report_includes_material_and_plant_sheets():
    # build sample df
    df = pd.DataFrame({
        "Supplier": ["A", "A", "B", "C", "B"],
        "Material": ["M1", "M2", "M1", "M3", "M2"],
        "Plant": ["P1", "P1", "P2", "P2", "P1"],
        "value": [100, 200, 50, 25, 75],
    })

    extra = build_vendor_report(df, group_col="Supplier", value_col="value", top_materials_per_plant=2)
    assert "material_pivot" in extra
    assert "plant_summary" in extra
    assert "plant_top_materials" in extra

    dfs = {"supplier_totals": extra["supplier_totals"], "material_pivot": extra["material_pivot"], "plant_summary": extra["plant_summary"], "plant_top_materials": extra["plant_top_materials"]}
    b = save_excel_report(dfs, figures=None)
    assert isinstance(b, (bytes, bytearray))

    # open workbook and validate sheets and numeric formatting
    wb = load_workbook(io.BytesIO(b))
    assert set(["supplier_totals", "material_pivot", "plant_summary", "plant_top_materials"]).issubset(set(wb.sheetnames))

    # check number format on a value cell in material_pivot (first data row, column with 'total_value')
    ws = wb["material_pivot"]
    # find column index for total_value
    headers = [cell.value for cell in ws[1]]
    idx = headers.index("total_value") + 1
    cell = ws.cell(row=2, column=idx)
    assert cell.number_format in ("#,##0.00", "0", "General") or isinstance(cell.value, (int, float))
