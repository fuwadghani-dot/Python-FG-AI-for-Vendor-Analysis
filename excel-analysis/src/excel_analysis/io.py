"""I/O helpers for Excel files"""
from typing import Optional, Union
import pandas as pd
from io import BytesIO


def load_excel(path_or_buffer, sheet_name: Optional[Union[int, str]] = 0, dtype: Optional[dict] = None) -> pd.DataFrame:
    """Load an Excel file (path or file-like) into a DataFrame.

    Args:
        path_or_buffer: file path or file-like object
        sheet_name: sheet index or name
        dtype: optional dtype mapping for pandas

    Returns:
        DataFrame
    """
    return pd.read_excel(path_or_buffer, sheet_name=sheet_name, engine="openpyxl", dtype=dtype)


def save_excel(df: pd.DataFrame, path: str, sheet_name: str = "Sheet1") -> None:
    """Save DataFrame to an Excel file path."""
    df.to_excel(path, index=False, sheet_name=sheet_name, engine="openpyxl")


def to_excel_bytes(df: pd.DataFrame) -> bytes:
    """Return an in-memory Excel file as bytes (useful for downloads)."""
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, index=False)
    buf.seek(0)
    return buf.getvalue()


from typing import Optional

def save_excel_report(dfs: dict, figures: Optional[dict] = None) -> bytes:
    """Create an Excel report workbook containing provided DataFrames and embed figures as PNG images.

    Args:
        dfs: mapping of sheet name -> DataFrame
        figures: optional mapping of sheet name -> plotly Figure (embedded into that sheet at H2)

    Returns:
        bytes of the Excel workbook
    """
    from io import BytesIO
    import plotly.io as pio
    from openpyxl import load_workbook
    from openpyxl.drawing.image import Image as OpenpyxlImage
    from PIL import Image as PILImage

    buf = BytesIO()
    # write sheets
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        for name, df in dfs.items():
            safe_name = str(name)[:31]
            df.to_excel(writer, sheet_name=safe_name, index=False)
    buf.seek(0)

    # embed images if provided
    wb = load_workbook(buf)
    if figures:
        for sheet_name, fig in figures.items():
            if fig is None:
                continue
            try:
                img_bytes = pio.to_image(fig, format="png", engine="kaleido")
            except Exception:
                # fallback: skip embedding this figure
                continue
            img_buf = BytesIO(img_bytes)
            pil = PILImage.open(img_buf)
            img = OpenpyxlImage(pil)
            target = str(sheet_name)[:31]
            if target not in wb.sheetnames:
                target = wb.sheetnames[0]
            ws = wb[target]
            # place image at H2
            ws.add_image(img, "H2")

    # Apply basic styling: column widths and number formats based on DataFrame content
    from openpyxl.utils import get_column_letter

    for sheet_name, df in dfs.items():
        name = str(sheet_name)[:31]
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        # set column widths based on header and sample values
        for i, col in enumerate(df.columns, start=1):
            col_letter = get_column_letter(i)
            header = str(col)
            max_len = len(header) + 2
            # check first 200 rows for width
            for val in df[col].astype(str).head(200):
                if val is None:
                    continue
                max_len = max(max_len, len(val) + 1)
            ws.column_dimensions[col_letter].width = max_len

            # set number format for numeric columns
            if pd.api.types.is_integer_dtype(df[col]) or pd.api.types.is_float_dtype(df[col]):
                # use two decimal places for floats, integer format for ints
                fmt = "0" if pd.api.types.is_integer_dtype(df[col]) else "#,##0.00"
                # apply to the column cells (skip header row)
                for row_idx in range(2, 2 + len(df.index)):
                    cell = ws[f"{col_letter}{row_idx}"]
                    cell.number_format = fmt
            # percentage column heuristic
            if header.lower().find("percent") >= 0 or header.lower().find("percentage") >= 0 or header.lower().endswith("%"):
                for row_idx in range(2, 2 + len(df.index)):
                    cell = ws[f"{col_letter}{row_idx}"]
                    cell.number_format = "0.00%"

    out = BytesIO()
    wb.save(out)
    out.seek(0)
    return out.getvalue()
